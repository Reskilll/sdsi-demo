"""Build the editable PowerPoint deck: presentation/AI-EDR_Guardrail.pptx

    python presentation/build_pptx.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from content import ROOT, SCREENS, slides  # noqa: E402

# White theme
BG = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF4, 0xF6, 0xFA)
LINE = RGBColor(0xDC, 0xE2, 0xEC)
TEXT = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x5B, 0x64, 0x78)
CYAN = RGBColor(0x0B, 0x6B, 0xCB)        # primary accent (blue)
BLUE = RGBColor(0x25, 0x63, 0xEB)
TEAM = "checkpoint_zero"
TERM_TEXT = RGBColor(0xE6, 0xEB, 0xF5)
RED = RGBColor(0xEF, 0x44, 0x44)
FONT = "Segoe UI"
MONO = "Menlo"
W, H = Inches(13.333), Inches(7.5)


def text(slide, x, y, w, h, s, size=18, color=TEXT, bold=False, font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = s if isinstance(s, list) else [s]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.name = font
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    s.shadow.inherit = False
    return s


def base(prs, title: str | None, n: int, total: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    box(slide, 0, 0, W, Inches(0.08), fill=CYAN, line=None, shape=MSO_SHAPE.RECTANGLE)
    if title:
        text(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.9), title, size=32, bold=True)
    text(slide, Inches(0.6), Inches(7.0), Inches(8), Inches(0.35), f"{TEAM} · AI-EDR Inline Guardrail · Digithon 2026", size=10, color=MUTED)
    text(slide, Inches(11.7), Inches(7.0), Inches(1.0), Inches(0.35), f"{n} / {total}", size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def fit_image(slide, path: str, x, y, w, h):
    iw, ih = Image.open(path).size
    scale = min(w / iw, h / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    return slide.shapes.add_picture(path, x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)


def speaker_notes() -> dict[int, str]:
    """Presenter-view notes from docs/SPEAKER_SCRIPT.md ("## Slide N — ..." sections)."""
    import re
    path = ROOT / "docs" / "SPEAKER_SCRIPT.md"
    if not path.exists():
        return {}
    notes = {}
    for block in re.split(r"^## ", path.read_text(), flags=re.M)[1:]:
        m = re.match(r"Slide (\d+)[^\n]*\n", block)
        if not m:
            continue
        body = block[m.end():].split("\n---")[0]
        body = re.sub(r"\*\*(.+?)\*\*", r"\1", body).replace("*(", "(").replace(")*", ")")
        notes[int(m.group(1))] = re.sub(r"\n{3,}", "\n\n", body).strip()
    return notes


def build(out: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    deck = slides()
    script = speaker_notes()
    for i, s in enumerate(deck, 1):
        kind = s["kind"]
        if kind == "team":
            slide = base(prs, None, i, len(deck))
            box(slide, 0, 0, Inches(0.35), H, fill=CYAN, line=None, shape=MSO_SHAPE.RECTANGLE)
            text(slide, Inches(1.0), Inches(0.9), Inches(11), Inches(0.5), "THEME", size=14, bold=True, color=MUTED)
            text(slide, Inches(1.0), Inches(1.3), Inches(11), Inches(0.7), s["theme"], size=30, bold=True, color=CYAN)
            text(slide, Inches(1.0), Inches(2.3), Inches(11), Inches(0.5), "TEAM", size=14, bold=True, color=MUTED)
            text(slide, Inches(1.0), Inches(2.65), Inches(11.5), Inches(1.2), s["title"], size=60, bold=True, font=MONO)
            text(slide, Inches(1.0), Inches(3.85), Inches(11), Inches(0.5), s["project"], size=18, color=MUTED)
            members = s["members"]
            cw_ = Inches(2.25)
            for k, name in enumerate(members):
                x = Inches(1.0) + k * (cw_ + Inches(0.12))
                box(slide, x, Inches(4.85), cw_, Inches(1.45))
                initials = "".join(w[0] for w in name.split()[:2]).upper()
                c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.2), Inches(5.08), Inches(0.55), Inches(0.55))
                c.fill.solid(); c.fill.fore_color.rgb = CYAN; c.line.fill.background()
                text(slide, x + Inches(0.2), Inches(5.17), Inches(0.55), Inches(0.4), initials, size=13, bold=True,
                     color=BG, align=PP_ALIGN.CENTER)
                text(slide, x + Inches(0.15), Inches(5.72), cw_ - Inches(0.25), Inches(0.5), name, size=14, bold=True)
        elif kind == "title":
            slide = base(prs, None, i, len(deck))
            box(slide, Inches(0.6), Inches(2.0), Inches(0.9), Inches(0.9), fill=CYAN, line=None)
            text(slide, Inches(0.6), Inches(2.08), Inches(0.9), Inches(0.75), "EDR", size=20, bold=True, color=BG, align=PP_ALIGN.CENTER)
            text(slide, Inches(0.6), Inches(3.1), Inches(12), Inches(1.2), s["title"], size=54, bold=True)
            text(slide, Inches(0.6), Inches(4.25), Inches(12), Inches(0.8), s["subtitle"], size=24, color=CYAN)
            text(slide, Inches(0.6), Inches(5.2), Inches(12), Inches(0.6), s["foot"], size=16, color=MUTED)
        else:
            slide = base(prs, s["title"], i, len(deck))
        top, left, cw, ch = Inches(1.45), Inches(0.6), Inches(12.13), Inches(5.35)

        if kind == "image":
            fit_image(slide, s["image"], left, top, cw, ch - Inches(0.6))
            text(slide, left, top + ch - Inches(0.5), cw, Inches(0.5), s["caption"], size=14, color=MUTED, align=PP_ALIGN.CENTER)
        elif kind == "video":
            mp4 = HERE / "build" / f"{s['video']}.mp4"
            poster = Path(s.get("poster") or SCREENS / "playground_fabrix.png")
            if mp4.exists():
                iw, ih = Image.open(poster).size
                scale = min(cw / iw, (ch - Inches(0.3)) / ih)
                pw, ph = int(iw * scale), int(ih * scale)
                slide.shapes.add_movie(str(mp4), left + (cw - pw) // 2, top, pw, ph, poster_frame_image=str(poster),
                                       mime_type="video/mp4")
            else:
                fit_image(slide, str(poster), left, top, cw, ch - Inches(0.3))
            text(slide, left, top + ch - Inches(0.3), cw, Inches(0.4), s.get("caption") or "Click to play · or run the live console at http://localhost:8080/console",
                 size=12, color=MUTED, align=PP_ALIGN.CENTER)
        elif kind == "cards":
            cards = s["cards"]
            cols = 2
            gw, gh = (cw - Inches(0.3)) // cols, (ch - Inches(0.3)) // 2
            for k, (t, d) in enumerate(cards):
                x = left + (k % cols) * (gw + Inches(0.3))
                y = top + (k // cols) * (gh + Inches(0.3))
                box(slide, x, y, gw, gh)
                box(slide, x, y + Inches(0.25), Inches(0.08), Inches(0.5), fill=CYAN, line=None, shape=MSO_SHAPE.RECTANGLE)
                text(slide, x + Inches(0.35), y + Inches(0.22), gw - Inches(0.6), Inches(0.6), t, size=22, bold=True)
                text(slide, x + Inches(0.35), y + Inches(0.95), gw - Inches(0.6), gh - Inches(1.1), d, size=16, color=MUTED)
        elif kind == "flow":
            steps = s["steps"]
            gap = Inches(0.45)
            bw = (cw - gap * (len(steps) - 1)) // len(steps)
            for k, (n, t, d) in enumerate(steps):
                x = left + k * (bw + gap)
                box(slide, x, top + Inches(0.3), bw, Inches(3.6))
                text(slide, x + Inches(0.25), top + Inches(0.5), bw, Inches(0.4), f"STAGE {n}", size=13, bold=True, color=CYAN)
                text(slide, x + Inches(0.25), top + Inches(0.95), bw - Inches(0.5), Inches(1.0), t, size=22, bold=True)
                text(slide, x + Inches(0.25), top + Inches(2.0), bw - Inches(0.5), Inches(1.8), d, size=15, color=MUTED)
                if k < len(steps) - 1:
                    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + bw + Inches(0.07), top + Inches(1.9), Inches(0.31), Inches(0.35))
                    a.fill.solid()
                    a.fill.fore_color.rgb = BLUE
                    a.line.fill.background()
            text(slide, left, top + Inches(4.3), cw, Inches(0.6), s["note"], size=16, color=MUTED, align=PP_ALIGN.CENTER)
        elif kind == "table":
            rows = [s["header"]] + list(s["rows"])
            shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top + Inches(0.1), cw, Inches(0.7) * len(rows))
            tbl = shape.table
            widths = list(s.get("widths") or (0.24, 0.56, 0.20))
            for c, wf in enumerate(widths[:len(rows[0])]):
                tbl.columns[c].width = Emu(int(cw * wf))
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = tbl.cell(r, c)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(0xE8, 0xEE, 0xF8) if r == 0 else (BG if r % 2 else PANEL)
                    cell.margin_left = cell.margin_right = Inches(0.15)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    tf = cell.text_frame
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.text = val
                    for run in p.runs:
                        run.font.size = Pt(13 if r else 13)
                        run.font.bold = r == 0 or c == 0
                        run.font.color.rgb = MUTED if r == 0 else (TEXT if c != 2 else CYAN)
                        run.font.name = MONO if (c == 1 and "/" in val) else FONT
        elif kind == "terminal":
            box(slide, left, top, cw, ch - Inches(0.1), fill=RGBColor(0x07, 0x0B, 0x16))
            for k, c in enumerate((RED, RGBColor(0xF5, 0x9E, 0x0B), RGBColor(0x22, 0xC5, 0x5E))):
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.25 + 0.25 * k), top + Inches(0.2), Inches(0.14), Inches(0.14))
                dot.fill.solid()
                dot.fill.fore_color.rgb = c
                dot.line.fill.background()
            tb = text(slide, left + Inches(0.3), top + Inches(0.55), cw - Inches(0.6), ch - Inches(0.8), "", size=14, font=MONO)
            tf = tb.text_frame
            for k, line in enumerate(s["lines"]):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                r = p.add_run()
                r.text = line or " "
                r.font.name = MONO
                r.font.size = Pt(14)
                r.font.color.rgb = (CYAN if line.startswith("$") else RED if "Blocked" in line or "Error" in line
                                    else RGBColor(0xFC, 0xD3, 0x4D) if line.startswith((" -", "Suspicious")) else TERM_TEXT)
        elif kind == "metrics":
            ms = s["metrics"]
            cols = 3
            gw, gh = (cw - Inches(0.3) * (cols - 1)) // cols, (ch - Inches(0.3)) // 2
            for k, (v, label, sub) in enumerate(ms):
                x = left + (k % cols) * (gw + Inches(0.3))
                y = top + (k // cols) * (gh + Inches(0.3))
                box(slide, x, y, gw, gh)
                text(slide, x + Inches(0.3), y + Inches(0.3), gw, Inches(1.0), v, size=44, bold=True, color=CYAN)
                text(slide, x + Inches(0.3), y + Inches(1.35), gw - Inches(0.5), Inches(0.5), label, size=18, bold=True)
                text(slide, x + Inches(0.3), y + Inches(1.85), gw - Inches(0.5), Inches(0.6), sub, size=13, color=MUTED)
        slide.notes_slide.notes_text_frame.text = script.get(i, s["narration"])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build(HERE / "AI-EDR_Guardrail.pptx"))
